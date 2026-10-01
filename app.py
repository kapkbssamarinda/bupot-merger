import os
import io
import re
from datetime import datetime
from flask import Flask, render_template, request, jsonify, send_file
import pdfplumber
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
template_dir = os.path.join(BASE_DIR, 'templates')
if not os.path.isdir(template_dir):
    template_dir = os.path.join(os.path.dirname(BASE_DIR), 'templates')

app = Flask(__name__, template_folder=template_dir)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max upload

MONTHS_ID = {
    '01': 'Januari', '1': 'Januari',
    '02': 'Februari', '2': 'Februari',
    '03': 'Maret', '3': 'Maret',
    '04': 'April', '4': 'April',
    '05': 'Mei', '5': 'Mei',
    '06': 'Juni', '6': 'Juni',
    '07': 'Juli', '7': 'Juli',
    '08': 'Agustus', '8': 'Agustus',
    '09': 'September', '9': 'September',
    '10': 'Oktober',
    '11': 'November',
    '12': 'Desember'
}

def parse_v2_date(date_str):
    if not date_str:
        return ''
    m = re.search(r'(\d[\d\s]*?)\s*dd\s*(\d[\d\s]*?)\s*mm\s*(\d[\d\s]*?)\s*yyyy', date_str, re.IGNORECASE)
    if m:
        d = re.sub(r'\s+', '', m.group(1)).zfill(2)
        mo = re.sub(r'\s+', '', m.group(2)).zfill(2)
        y = re.sub(r'\s+', '', m.group(3))
        mo_name = MONTHS_ID.get(mo, mo)
        return f"{d} {mo_name} {y}"
    m2 = re.search(r'(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})', date_str)
    if m2:
        return m2.group(0).strip()
    return date_str.strip()

def parse_id_number(val, is_rate=False):
    if not val:
        return 0.0
    val_str = str(val).strip()
    if is_rate:
        val_clean = val_str.replace(',', '.')
        try:
            return float(val_clean)
        except ValueError:
            return 0.0
    if ',' in val_str:
        val_clean = val_str.replace('.', '').replace(',', '.')
    else:
        parts = val_str.split('.')
        if len(parts) > 2:
            val_clean = val_str.replace('.', '')
        elif len(parts) == 2 and len(parts[1]) == 3:
            val_clean = val_str.replace('.', '')
        else:
            val_clean = val_str
    try:
        return float(val_clean)
    except ValueError:
        return 0.0

def parse_bupot_text(text, source_name=""):
    """
    Ekstraksi data bukti potong dari teks PDF dengan deteksi otomatis versi:
    - V1 (Standar): Format DJP BPPU Unifikasi Berformat Standar
    - V2 (BPBS): Formulir BPBS PPh Pasal 4(2), 15, 22, dan 23
    - V3 (BPBS): Formulir BPBS Format Baru 16 Digit/NITKU (NPWP 15-digit terpisah, C.3 Nama, C.4 Tanggal)

    Kolom target utama:
    - NPWP PEMOTONG
    - NAMA PEMOTONG
    - TANGGAL PEMOTONGAN
    - DPP
    - PPH DIPOTONG
    - NOMOR BUPOT
    """
    is_bpbs = ('FORMULIR BPBS' in text) or ('H.1 NOMOR' in text)

    if is_bpbs:
        # Deteksi versi: V3 jika terdapat NITKU atau format NPWP ganda dengan '/' di C.1
        is_v3 = ('NITKU' in text) or bool(re.search(r'C\.1\s*[:\s]*NPWP[^\n\r]*\/', text))
        version = 'V3 (BPBS)' if is_v3 else 'V2 (BPBS)'

        # 1. Nomor Bupot: H.1 NOMOR : 2 0 0 0 0 1 3 5 6 5
        m_no = re.search(r'H\.1\s+NOMOR\s*:\s*([0-9\s]+?)(?=\s*H\.[0-9]|\n|$)', text)
        nomor = re.sub(r'\s+', '', m_no.group(1)) if m_no else ''

        # 2. Sifat Pajak: H.4 PPh Final vs H.5 X PPh Tidak Final
        sifat = 'TIDAK FINAL'
        if re.search(r'H\.4\s*[XxVv]\s*PPh\s*Final', text) or re.search(r'H\.4\s*\[[XxVv]\]\s*PPh\s*Final', text):
            sifat = 'FINAL'
        elif re.search(r'H\.5\s*[XxVv]\s*PPh\s*Tidak\s*Final', text) or re.search(r'H\.5\s*\[[XxVv]\]\s*PPh\s*Tidak\s*Final', text):
            sifat = 'TIDAK FINAL'

        # 3. Status: H.2 Pembetulan Ke- 0
        m_stat = re.search(r'H\.2\s+Pembetulan\s+Ke-\s*(\d+)', text, re.IGNORECASE)
        status = f'PEMBETULAN-{m_stat.group(1)}' if m_stat and m_stat.group(1) != '0' else 'NORMAL'

        # 4. Table B.1 - B.6
        masa_pajak = ''
        kode_objek = ''
        dpp = 0.0
        tarif = 0.0
        pph = 0.0

        sec_b_m = re.search(r'B\.1\s+B\.2\s+B\.3\s+B\.4\s+B\.5\s+B\.6\s*([\s\S]*?)(?:Keterangan\s+Kode\s+Objek|B\.7)', text)
        if sec_b_m:
            lines = [l.strip() for l in sec_b_m.group(1).splitlines() if l.strip()]
            if lines:
                parts = lines[0].split()
                if len(parts) >= 5:
                    raw_masa = parts[0]
                    if '-' in raw_masa:
                        m_parts = raw_masa.split('-')
                        masa_pajak = f"{m_parts[0].zfill(2)}-{m_parts[1]}"
                    else:
                        masa_pajak = raw_masa
                    kode_objek = parts[1]
                    dpp = parse_id_number(parts[2])
                    if len(parts) == 5:
                        tarif = parse_id_number(parts[3], is_rate=True)
                        pph = parse_id_number(parts[4])
                    else:
                        tarif = parse_id_number(parts[4], is_rate=True)
                        pph = parse_id_number(parts[5])

        # 5. Jenis PPh derived from kode_objek or form
        if kode_objek.startswith('24-'):
            jenis_pph = 'Pasal 23'
        elif kode_objek.startswith('22-'):
            jenis_pph = 'Pasal 22'
        elif kode_objek.startswith('28-'):
            jenis_pph = 'Pasal 4 ayat (2)'
        elif kode_objek.startswith('15-'):
            jenis_pph = 'Pasal 15'
        else:
            jenis_pph = 'PPh Unifikasi'

        # 6. Objek Pajak description
        m_obj = re.search(r'Keterangan\s+Kode\s+Objek\s+Pajak\s*:\s*([\s\S]*?)B\.7', text)
        objek_pajak = ' '.join(m_obj.group(1).split()) if m_obj else ''

        # 7. Dokumen Referensi (B.7)
        m_dok = re.search(r'B\.7\s+Dokumen\s+Referensi[\s\S]*?Nomor\s+Dokumen\s*[:\s]\s*([^\n\r]+)', text)
        no_dokumen = m_dok.group(1).strip() if m_dok else ''

        m_tgldok = re.search(r'B\.7\s+Dokumen\s+Referensi[\s\S]*?Tanggal\s*[:\s]\s*([0-9\s]+dd[0-9\s]+mm[0-9\s]+yyyy|[0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})', text)
        tgl_dokumen = parse_v2_date(m_tgldok.group(1)) if m_tgldok else ''

        # 8. NPWP Pemotong (C.1) - Ambil 15 digit pertama saja (sebelum '/' jika format baru)
        m_npwp = re.search(r'C\.1\s*[:\s]*NPWP\s*[:\s]*([^\n\r]+)', text)
        if m_npwp:
            raw_npwp = m_npwp.group(1).split('/')[0]
            npwp = re.sub(r'\D', '', raw_npwp)[:15]
        else:
            npwp = ''

        # 9. Nama Pemotong (C.2 pada V2 atau C.3 pada V3)
        m_nama = re.search(r'C\.[0-9]\s+Nama\s+(?:Wajib\s+Pajak|Pemotong)[^\n:]*:\s*([^\n\r]+)', text)
        nama_pemotong = m_nama.group(1).strip() if m_nama else ''

        # 10. Tanggal Pemotongan (C.3 pada V2 atau C.4 pada V3)
        m_tglpotong = re.search(r'C\.[0-9]\s+Tanggal\s*:\s*([0-9\s]+dd[0-9\s]+mm[0-9\s]+yyyy|[0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4}|[^\n\r]+)', text)
        tgl_pemotongan = parse_v2_date(m_tglpotong.group(1)) if m_tglpotong else ''

    else:
        version = 'V1 (Standar)'

        # 1. Header row: NOMOR, MASA PAJAK, SIFAT, STATUS
        header_m = re.search(r'([A-Z0-9]{8,16})\s+(\d{2}-\d{4})\s+([A-Z]+)\s+([A-Z]+)', text)
        nomor = header_m.group(1) if header_m else ''
        masa_pajak = header_m.group(2) if header_m else ''
        sifat = header_m.group(3) if header_m else ''
        status = header_m.group(4) if header_m else ''

        # 2. Jenis PPh: B.2 Jenis PPh : Pasal 22
        jenis_pph_m = re.search(r'B\.2\s+Jenis\s+PPh\s*:\s*([^\n\r]+)', text)
        jenis_pph = jenis_pph_m.group(1).strip() if jenis_pph_m else ''

        # 3. Kode Objek Pajak, Objek Pajak, DPP, Tarif, Pajak Penghasilan (B.3 - B.7)
        sec_b = re.search(r'B\.3\s+B\.4\s+B\.5\s+B\.6\s+B\.7\s*([\s\S]*?)B\.8', text)
        kode_objek = ''
        objek_pajak = ''
        dpp = 0.0
        tarif = 0.0
        pph = 0.0

        if sec_b:
            b_text = sec_b.group(1).strip()
            lines = [l.strip() for l in b_text.splitlines() if l.strip()]
            if lines:
                m_row = re.search(r'^(\d{2}-\d{3}-\d{2})\s+(.*?)\s+([\d\.,]+)\s+([\d\.,]+)\s+([\d\.,]+)$', lines[0])
                if m_row:
                    kode_objek = m_row.group(1)
                    objek_p1 = m_row.group(2)
                    dpp = parse_id_number(m_row.group(3))
                    tarif = parse_id_number(m_row.group(4), is_rate=True)
                    pph = parse_id_number(m_row.group(5))
                    remaining_desc = ' '.join(lines[1:])
                    objek_pajak = (objek_p1 + (' ' + remaining_desc if remaining_desc else '')).strip()
                else:
                    objek_pajak = ' '.join(lines)

        # 4. Tanggal Dokumen: B.8 Dokumen Dasar Bukti ... Tanggal : 07 September 2026
        tgl_dok_m = re.search(r'B\.8[\s\S]*?Tanggal\s*:\s*([0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})', text)
        tgl_dokumen = tgl_dok_m.group(1).strip() if tgl_dok_m else ''

        # 5. Nomor Dokumen: B.9 Nomor Dokumen : 9140424520
        no_dok_m = re.search(r'B\.9\s+Nomor\s+Dokumen\s*:\s*([^\n\r]+)', text)
        no_dokumen = no_dok_m.group(1).strip() if no_dok_m else ''

        # 6. NPWP Pemotong: C.1 NPWP / NIK : 0010611572051000
        npwp_m = re.search(r'C\.1\s+NPWP\s*/\s*NIK\s*:\s*([0-9]{15,16})', text)
        npwp = npwp_m.group(1).strip() if npwp_m else ''

        # 7. Nama Pemotong: C.3 NAMA PEMOTONG DAN/ATAU PEMUNGUT PPh : PERTAMINA PATRA NIAGA
        nama_pemotong_m = re.search(r'C\.3\s+NAMA\s+PEMOTONG[^\n:]*:\s*([^\n\r]+)', text)
        nama_pemotong = nama_pemotong_m.group(1).strip() if nama_pemotong_m else ''

        # 8. Tanggal Pemotongan: C.4 TANGGAL : 07 September 2026
        tgl_potong_m = re.search(r'C\.4\s+TANGGAL\s*:\s*([0-9]{1,2}\s+[A-Za-z]+\s+[0-9]{4})', text)
        tgl_pemotongan = tgl_potong_m.group(1).strip() if tgl_potong_m else ''

    return {
        'file_source': source_name,
        'versi': version,
        'Nomor': nomor,
        'Masa Pajak': masa_pajak,
        'Sifat Pajak Penghasilan': sifat,
        'Status': status,
        'jenis pph': jenis_pph,
        'kode objek pajak': kode_objek,
        'objek pajak': objek_pajak,
        'dpp': dpp,
        'tarif': tarif,
        'pajak penghasilan': pph,
        'tanggal dokumen': tgl_dokumen,
        'nomor dokumen': no_dokumen,
        'npwp': npwp,
        'nama pemotong': nama_pemotong,
        'tanggal pemotongan': tgl_pemotongan
    }

def extract_pdf_file(file_bytes_or_path, filename=""):
    try:
        with pdfplumber.open(file_bytes_or_path) as pdf:
            full_text = "\n".join([page.extract_text() or "" for page in pdf.pages])
        return parse_bupot_text(full_text, source_name=filename)
    except Exception as e:
        return {
            'file_source': filename,
            'versi': 'Tidak Dikenal',
            'error': str(e),
            'Nomor': '',
            'Masa Pajak': '',
            'Sifat Pajak Penghasilan': '',
            'Status': 'GAGAL',
            'jenis pph': '',
            'kode objek pajak': '',
            'objek pajak': f'Gagal membaca file: {e}',
            'dpp': 0.0,
            'tarif': 0.0,
            'pajak penghasilan': 0.0,
            'tanggal dokumen': '',
            'nomor dokumen': '',
            'npwp': '',
            'nama pemotong': '',
            'tanggal pemotongan': ''
        }

@app.route('/')
@app.route('/api')
@app.route('/api/index')
@app.route('/api/index.py')
def index():
    return render_template('index.html')

import urllib.parse

class VercelPathMiddleware:
    def __init__(self, wsgi_app):
        self.wsgi_app = wsgi_app

    def __call__(self, environ, start_response):
        qs = environ.get('QUERY_STRING', '')
        if '__route__' in qs:
            params = urllib.parse.parse_qs(qs)
            if '__route__' in params and params['__route__']:
                route = params['__route__'][0]
                if not route.startswith('/'):
                    route = '/' + route
                environ['PATH_INFO'] = route
                # Bersihkan parameter __route__ dari QUERY_STRING
                filtered = [(k, v) for k, vs in params.items() if k != '__route__' for v in vs]
                environ['QUERY_STRING'] = urllib.parse.urlencode(filtered)
        else:
            path = environ.get('PATH_INFO', '')
            if path in ('/api/index.py', '/api/index', '/api', ''):
                environ['PATH_INFO'] = '/'
            elif path.startswith('/api/index.py/'):
                environ['PATH_INFO'] = path[len('/api/index.py'):]
        return self.wsgi_app(environ, start_response)

app.wsgi_app = VercelPathMiddleware(app.wsgi_app)

@app.route('/api/extract-upload', methods=['POST'])
def extract_from_upload():
    uploaded_files = request.files.getlist('files')
    if not uploaded_files or uploaded_files[0].filename == '':
        return jsonify({'success': False, 'message': 'Tidak ada file yang diunggah.'}), 400

    results = []
    for f in uploaded_files:
        if f.filename.lower().endswith('.pdf'):
            content = f.read()
            size_kb = f"{max(1, round(len(content) / 1024))} KB"
            file_stream = io.BytesIO(content)
            res = extract_pdf_file(file_stream, filename=f.filename)
            res['file_size'] = size_kb
            results.append(res)

    if not results:
        return jsonify({'success': False, 'message': 'Tidak ada file PDF valid yang diproses.'}), 400

    return jsonify({'success': True, 'data': results, 'total_files': len(results)})

@app.route('/api/export-excel', methods=['POST'])
def export_excel():
    data = request.json.get('data', [])
    if not data:
        return jsonify({'success': False, 'message': 'Data kosong, tidak dapat membuat Excel.'}), 400

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Bupot"

    # Kolom hasil konsolidasi mencakup Versi Bupot & 15 kolom standar:
    headers = [
        "VERSI",
        "NOMOR BUPOT",
        "MASA PAJAK",
        "SIFAT PAJAK PENGHASILAN",
        "STATUS",
        "JENIS PPH",
        "KODE OBJEK PAJAK",
        "OBJEK PAJAK",
        "DPP",
        "TARIF",
        "PPH DIPOTONG",
        "TANGGAL DOKUMEN",
        "NOMOR DOKUMEN",
        "NPWP PEMOTONG",
        "NAMA PEMOTONG",
        "TANGGAL PEMOTONGAN"
    ]

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Tulis Header
    ws.append(headers)
    ws.row_dimensions[1].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border

    # Tulis Data Baris
    row_start = 2
    for item in data:
        row_values = [
            str(item.get("versi", "")),
            str(item.get("Nomor", "")),
            str(item.get("Masa Pajak", "")),
            str(item.get("Sifat Pajak Penghasilan", "")),
            str(item.get("Status", "")),
            str(item.get("jenis pph", "")),
            str(item.get("kode objek pajak", "")),
            str(item.get("objek pajak", "")),
            float(item.get("dpp", 0.0) or 0.0),
            float(item.get("tarif", 0.0) or 0.0),
            float(item.get("pajak penghasilan", 0.0) or 0.0),
            str(item.get("tanggal dokumen", "")),
            str(item.get("nomor dokumen", "")),
            str(item.get("npwp", "")),
            str(item.get("nama pemotong", "")),
            str(item.get("tanggal pemotongan", ""))
        ]
        ws.append(row_values)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 20

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

            # Formatting khusus
            if col_idx in [1, 2, 3, 4, 5, 7, 12, 16]:  # VERSI, NOMOR, MASA, SIFAT, STATUS, KODE, TGL DOK, TGL POTONG
                cell.alignment = center_align
            elif col_idx in [9, 11]:  # DPP & PPh Dipotong (Numeric with thousand separator)
                cell.number_format = '#,##0'
                cell.alignment = right_align
            elif col_idx == 10:  # Tarif
                cell.number_format = '0.0%' if cell.value < 1 else '0.0'
                cell.alignment = right_align
            elif col_idx == 14:  # NPWP Pemotong (Text format preserving leading zero)
                cell.number_format = '@'
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Baris Total / Summary
    last_data_row = ws.max_row
    if last_data_row >= row_start:
        total_row = last_data_row + 1
        ws.cell(row=total_row, column=1, value="TOTAL")
        ws.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=8)

        # Formula SUM DPP (Kolom I / 9)
        ws.cell(row=total_row, column=9, value=f"=SUM(I{row_start}:I{last_data_row})")
        # Formula SUM PPh (Kolom K / 11)
        ws.cell(row=total_row, column=11, value=f"=SUM(K{row_start}:K{last_data_row})")

        total_font = Font(name="Calibri", size=10, bold=True, color="0F172A")
        total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        thick_top_double_bottom = Border(
            top=Side(style='thin', color='0F172A'),
            bottom=Side(style='double', color='0F172A'),
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1')
        )

        ws.row_dimensions[total_row].height = 22
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = thick_top_double_bottom
            if col_idx in [9, 11]:
                cell.number_format = '#,##0'
                cell.alignment = right_align
            elif col_idx == 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Auto-fit Column Widths dengan padding
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if cell.number_format == '#,##0' and isinstance(cell.value, (int, float)):
                val = f"{cell.value:,.0f}"
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Rekap_Merger_Bupot_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )

if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    print(f"Aplikasi Merger Bupot berjalan di http://localhost:{port}")
    app.run(host='0.0.0.0', port=port, debug=True)
